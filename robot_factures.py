"""
TP RPA - Robot de traitement automatique des factures
Étapes : détection -> extraction -> décision -> rapport Excel -> classement -> notification
Gestion d'erreurs + journal log.txt
"""
import logging
import os
import re
import shutil
import smtplib
from datetime import datetime
from email.message import EmailMessage
from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

# ------------------------------------------------------------------ CONFIG
DOSSIER = Path("Factures")
SOURCE_EXCEL = Path("factures.xlsx")   # fichier Excel de départ contenant les factures
DOSSIER_NORMALES = DOSSIER / "Traitees"
DOSSIER_CONTROLE = DOSSIER / "A_controler"
DOSSIER_ERREURS = DOSSIER / "Erreurs"
RAPPORT = Path("rapport_factures.xlsx")
SEUIL = 1000  # MAD

# Paramètres e-mail (variables d'environnement). Si absents -> mode simulation.
SMTP_HOST = os.getenv("SMTP_HOST", "smtp.gmail.com")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))
SMTP_USER = os.getenv("SMTP_USER")            # ex : votre.adresse@gmail.com
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD")    # mot de passe d'application
EMAIL_RESPONSABLE = os.getenv("EMAIL_RESPONSABLE", "responsable@exemple.ma")

# ------------------------------------------------------------------ LOG
logging.basicConfig(
    filename="log.txt",
    level=logging.INFO,
    format="%(asctime)s | %(levelname)-7s | %(message)s",
    datefmt="%d/%m/%Y %H:%M:%S",
    encoding="utf-8",
)
log = logging.getLogger("robot")


def info(msg):
    print(msg)
    log.info(msg)


def erreur(msg):
    print("  [ERREUR]", msg)
    log.error(msg)


# ------------------------------------------------------------------ EXCEPTIONS
class FactureError(Exception):
    """Erreur de traitement d'une facture."""


# ------------------------------------------------------------------ 2. LECTURE / EXTRACTION
COLONNES = ["N°", "Fournisseur", "Date", "Montant", "Fichier"]


def lire_excel_source(chemin: Path) -> pd.DataFrame:
    """Lit le fichier Excel contenant les factures. Lève FactureError s'il est illisible."""
    try:
        df = pd.read_excel(chemin, dtype=object)
    except Exception as e:
        raise FactureError(f"Fichier Excel illisible ({type(e).__name__})")
    manquantes = [c for c in COLONNES[:4] if c not in df.columns]
    if manquantes:
        raise FactureError("Colonnes manquantes dans l'Excel : " + ", ".join(manquantes))
    if "Fichier" not in df.columns:
        df["Fichier"] = ""
    return df


def vide(v) -> bool:
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == ""


def convertir_montant(brut) -> float:
    """'2 750,50 MAD' -> 2750.5 ; lève FactureError si incorrect."""
    if isinstance(brut, (int, float)):
        valeur = float(brut)
    else:
        nettoye = re.sub(r"(?i)MAD|DH", "", str(brut)).replace(" ", "").replace(",", ".")
        try:
            valeur = float(nettoye)
        except ValueError:
            raise FactureError(f"Montant incorrect : '{str(brut).strip()}'")
    if valeur <= 0:
        raise FactureError(f"Montant incorrect (<= 0) : {valeur}")
    return valeur


def formater_date(v) -> str:
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.strftime("%d/%m/%y")
    return str(v).strip()


def extraire_infos(ligne) -> dict:
    """Extrait et valide une ligne de l'Excel (une facture)."""
    manquants = [c for c in COLONNES[:4] if vide(ligne[c])]
    if manquants:
        raise FactureError("Information manquante : " + ", ".join(manquants))
    return {
        "N°": str(ligne["N°"]).strip(),
        "Fournisseur": str(ligne["Fournisseur"]).strip(),
        "Date": formater_date(ligne["Date"]),
        "Montant": convertir_montant(ligne["Montant"]),
    }


def fichier_lisible(chemin: Path) -> bool:
    """Vérifie qu'un PDF existe et s'ouvre. Lève FactureError s'il est illisible."""
    try:
        with pdfplumber.open(chemin) as pdf:
            _ = len(pdf.pages)
    except Exception as e:
        raise FactureError(f"Fichier PDF illisible ({type(e).__name__})")
    return True


# ------------------------------------------------------------------ 3. DÉCISION
def decider(montant: float) -> str:
    return "Normal" if montant <= SEUIL else "À contrôler"


# ------------------------------------------------------------------ 5. CLASSEMENT
def classer(chemin: Path, dossier_dest: Path) -> Path:
    dossier_dest.mkdir(parents=True, exist_ok=True)
    destination = dossier_dest / chemin.name
    shutil.move(str(chemin), str(destination))
    return destination


# ------------------------------------------------------------------ 6. NOTIFICATION
def notifier(numero, fournisseur, montant):
    sujet = f"Facture à contrôler : {numero}"
    corps = (
        "Bonjour,\n\n"
        "Une facture dépasse le seuil de 1 000 MAD et doit être contrôlée :\n"
        f"  - Numéro      : {numero}\n"
        f"  - Fournisseur : {fournisseur}\n"
        f"  - Montant     : {montant:,.2f} MAD\n\n"
        "Cordialement,\nRobot RPA Factures"
    )
    msg = EmailMessage()
    msg["Subject"], msg["To"] = sujet, EMAIL_RESPONSABLE
    msg["From"] = SMTP_USER or "robot-rpa@exemple.ma"
    msg.set_content(corps)

    if not (SMTP_USER and SMTP_PASSWORD):
        # Mode simulation : l'e-mail est écrit dans un fichier
        with open("emails_simules.txt", "a", encoding="utf-8") as f:
            f.write(f"--- {datetime.now():%d/%m/%Y %H:%M:%S} ---\n"
                    f"À : {EMAIL_RESPONSABLE}\nObjet : {sujet}\n\n{corps}\n\n")
        info(f"  E-mail SIMULÉ pour {numero} (SMTP non configuré) -> emails_simules.txt")
        return

    try:
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=20) as s:
            s.starttls()
            s.login(SMTP_USER, SMTP_PASSWORD)
            s.send_message(msg)
        info(f"  E-mail envoyé à {EMAIL_RESPONSABLE} pour la facture {numero}")
    except Exception as e:
        erreur(f"Échec d'envoi de l'e-mail pour {numero} : {e}")


# ------------------------------------------------------------------ 4. RAPPORT EXCEL
def generer_rapport(df: pd.DataFrame):
    sortie = df.copy()
    sortie.to_excel(RAPPORT, index=False, sheet_name="Factures")

    wb = load_workbook(RAPPORT)
    ws = wb.active
    entete = PatternFill("solid", fgColor="1F4E78")
    rouge = PatternFill("solid", fgColor="F8CBAD")
    jaune = PatternFill("solid", fgColor="FFE699")
    vert = PatternFill("solid", fgColor="C6E0B4")

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = entete
        cell.alignment = Alignment(horizontal="center")

    for row in ws.iter_rows(min_row=2):
        statut = str(row[4].value)
        couleur = vert if statut == "Normal" else jaune if statut == "À contrôler" else rouge
        row[4].fill = couleur
        if isinstance(row[3].value, (int, float)):
            row[3].number_format = '#,##0.00" MAD"'

    for col, largeur in zip("ABCDEF", (10, 22, 12, 18, 18, 45)):
        ws.column_dimensions[col].width = largeur
    wb.save(RAPPORT)
    info(f"Rapport Excel généré : {RAPPORT}")


# ------------------------------------------------------------------ MAIN
def main():
    info("=" * 50)
    info("Démarrage du robot RPA - traitement des factures")

    # 1) Lecture du fichier Excel de départ
    if not SOURCE_EXCEL.exists():
        erreur(f"Fichier '{SOURCE_EXCEL}' introuvable. Lancez d'abord generer_factures.py")
        return
    try:
        source = lire_excel_source(SOURCE_EXCEL)
    except FactureError as e:
        erreur(f"{SOURCE_EXCEL} : {e}")
        return
    info(f"{len(source)} facture(s) détectée(s) dans {SOURCE_EXCEL}")

    lignes = []
    for _, ligne in source.iterrows():
        ident = str(ligne["N°"]) if not vide(ligne["N°"]) else f"(ligne {_ + 2})"
        info(f"Traitement de la facture {ident}")
        nom_fichier = "" if vide(ligne["Fichier"]) else str(ligne["Fichier"]).strip()
        chemin = DOSSIER / nom_fichier if nom_fichier else None
        try:
            infos = extraire_infos(ligne)
            if chemin is not None and chemin.exists():
                fichier_lisible(chemin)          # lève FactureError si illisible
            statut = decider(infos["Montant"])
            infos["Statut"] = statut
            info(f"  Extrait : {infos['N°']} | {infos['Fournisseur']} | "
                 f"{infos['Date']} | {infos['Montant']:.2f} MAD -> {statut}")

            dest = DOSSIER_NORMALES if statut == "Normal" else DOSSIER_CONTROLE
            if chemin is not None and chemin.exists():
                classer(chemin, dest)
                info(f"  Fichier {nom_fichier} classé dans {dest}/")
            else:
                info(f"  Aucun fichier PDF associé : classement ignoré")

            if statut == "À contrôler":
                notifier(infos["N°"], infos["Fournisseur"], infos["Montant"])
            lignes.append({**infos, "Remarque": ""})

        except FactureError as e:
            erreur(f"{ident} : {e}")
            if chemin is not None and chemin.exists():
                classer(chemin, DOSSIER_ERREURS)
            lignes.append({"N°": ident, "Fournisseur": "" if vide(ligne["Fournisseur"]) else ligne["Fournisseur"],
                           "Date": "" if vide(ligne["Date"]) else formater_date(ligne["Date"]),
                           "Montant": None, "Statut": "Erreur", "Remarque": str(e)})
        except Exception as e:  # filet de sécurité : le robot ne doit jamais s'arrêter
            erreur(f"{ident} : erreur inattendue : {e}")
            lignes.append({"N°": ident, "Fournisseur": "", "Date": "",
                           "Montant": None, "Statut": "Erreur", "Remarque": str(e)})

    # DataFrame Pandas
    df = pd.DataFrame(lignes, columns=["N°", "Fournisseur", "Date", "Montant", "Statut", "Remarque"])
    print("\n", df.to_string(index=False), "\n")

    if not df.empty:
        generer_rapport(df)

    nb_ok = (df["Statut"] != "Erreur").sum() if not df.empty else 0
    info(f"Terminé : {nb_ok} facture(s) traitée(s), {len(df) - nb_ok} erreur(s)")


if __name__ == "__main__":
    main()
